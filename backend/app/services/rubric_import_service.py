"""Parse CSV, Word, and PDF files into RubricCategory[] for the rubric builder."""
import re
import os
import logging

logger = logging.getLogger(__name__)


def _slugify_key(text: str) -> str:
    """Convert category name to a dict key."""
    s = text.lower().strip()
    s = re.sub(r'[^\w\s]', '', s)
    return re.sub(r'\s+', '_', s)[:60]


def _find_col(header: list[str], variants: list[str]) -> int | None:
    """Find column index matching any of the variant names."""
    for i, h in enumerate(header):
        for v in variants:
            if v == h or v in h or h in v:
                return i
    return None


def _parse_points(text: str) -> float:
    """Extract a numeric point value from text like '7.5', '10 pts', etc."""
    match = re.search(r'(\d+(?:\.\d+)?)', str(text))
    return float(match.group(1)) if match else 0


def _parse_extracted_table(rows: list[list[str]]) -> list[dict]:
    """Parse a table (list of rows, each row a list of strings) into categories."""
    if not rows or len(rows) < 2:
        return []

    header = [h.lower().strip().replace(' ', '_') for h in rows[0]]

    name_idx = _find_col(header, ['name', 'category', 'category_name', 'criteria', 'component'])
    points_idx = _find_col(header, ['max_points', 'points', 'total_points', 'max', 'weight'])
    desc_idx = _find_col(header, ['description', 'desc', 'details'])
    excellent_idx = _find_col(header, ['excellent', 'outstanding', '90-100'])
    good_idx = _find_col(header, ['good', 'proficient', '75-89'])
    satisfactory_idx = _find_col(header, ['satisfactory', 'developing', '60-74', 'adequate'])
    needs_imp_idx = _find_col(header, ['needs_improvement', 'unsatisfactory', 'poor', 'beginning'])

    if name_idx is None:
        return []

    categories = []
    for row in rows[1:]:
        if len(row) <= name_idx:
            continue
        cat_name = row[name_idx].strip()
        if not cat_name:
            continue

        def _safe_get(idx, row_data):
            if idx is not None and idx < len(row_data):
                return row_data[idx].strip()
            return ''

        categories.append({
            'key': _slugify_key(cat_name),
            'name': cat_name,
            'max_points': _parse_points(_safe_get(points_idx, row)) if points_idx is not None else 0,
            'description': _safe_get(desc_idx, row),
            'criteria': {
                'excellent': _safe_get(excellent_idx, row),
                'good': _safe_get(good_idx, row),
                'satisfactory': _safe_get(satisfactory_idx, row),
                'needs_improvement': _safe_get(needs_imp_idx, row),
            },
        })

    return categories


def _extract_sections_from_lines(lines: list[str]) -> list[dict]:
    """Extract rubric categories from unstructured text lines using heuristics."""
    categories = []
    # Pattern: "Name (X points)" or "Name - X pts" with optional leading number
    section_pattern = re.compile(
        r'^(?:\d+[\.\)]\s*)?'
        r'(.+?)'
        r'\s*[\(\-\:]\s*'
        r'(\d+(?:\.\d+)?)\s*(?:pts?|points?|marks?)?'
        r'\s*[\)\:]?'
        r'\s*(.*)?$',
        re.IGNORECASE,
    )

    i = 0
    while i < len(lines):
        match = section_pattern.match(lines[i])
        if match:
            name = match.group(1).strip().rstrip(':-(')
            points = float(match.group(2))
            desc = (match.group(3) or '').strip()

            # Collect subsequent non-section lines as description
            i += 1
            while i < len(lines) and not section_pattern.match(lines[i]):
                if not desc:
                    desc = lines[i]
                i += 1

            categories.append({
                'key': _slugify_key(name),
                'name': name,
                'max_points': points,
                'description': desc,
                'criteria': {
                    'excellent': '',
                    'good': '',
                    'satisfactory': '',
                    'needs_improvement': '',
                },
            })
        else:
            i += 1

    return categories


# ---- Public API ----

def parse_rubric_csv(file_path: str) -> list[dict]:
    """Parse a CSV file into RubricCategory[].

    Expected columns (case-insensitive, flexible naming):
      - name / category / category_name
      - max_points / points
      - description (optional)
      - excellent, good, satisfactory, needs_improvement (optional)
    """
    import pandas as pd

    df = pd.read_csv(file_path, sep=None, engine='python')
    df.columns = [c.strip().lower().replace(' ', '_') for c in df.columns]

    name_variants = ['name', 'category', 'category_name', 'rubric_category']
    points_variants = ['max_points', 'points', 'total_points', 'max']
    desc_variants = ['description', 'desc', 'details']

    name_col = next((c for c in name_variants if c in df.columns), None)
    points_col = next((c for c in points_variants if c in df.columns), None)
    desc_col = next((c for c in desc_variants if c in df.columns), None)

    if not name_col:
        raise ValueError(
            'CSV must have a column named "name", "category", or "category_name".'
        )

    categories = []
    for _, row in df.iterrows():
        cat_name = str(row[name_col]).strip()
        if not cat_name or cat_name.lower() == 'nan':
            continue

        cat = {
            'key': _slugify_key(cat_name),
            'name': cat_name,
            'max_points': float(row[points_col]) if points_col and pd.notna(row[points_col]) else 0,
            'description': str(row[desc_col]).strip() if desc_col and pd.notna(row[desc_col]) else '',
            'criteria': {
                'excellent': '',
                'good': '',
                'satisfactory': '',
                'needs_improvement': '',
            },
        }

        for level in ['excellent', 'good', 'satisfactory', 'needs_improvement']:
            if level in df.columns and pd.notna(row[level]):
                cat['criteria'][level] = str(row[level]).strip()

        categories.append(cat)

    if not categories:
        raise ValueError('No rubric categories found in CSV. Check column names.')

    return categories


def parse_rubric_docx(file_path: str) -> list[dict]:
    """Extract rubric categories from a Word document.

    Strategy 1: Parse tables (most rubrics are tabular).
    Strategy 2: Fall back to section-based body text parsing.
    """
    import docx

    doc = docx.Document(file_path)
    categories = []

    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        cats = _parse_extracted_table(rows)
        if cats:
            categories.extend(cats)

    if categories:
        return categories

    # Fallback: parse body text
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return _extract_sections_from_lines(lines)


def parse_rubric_pdf(file_path: str) -> list[dict]:
    """Extract rubric categories from a PDF.

    Strategy 1: PyMuPDF table extraction.
    Strategy 2: Fall back to full text parsing.
    """
    import fitz

    doc = fitz.open(file_path)
    categories = []

    for page in doc:
        try:
            tables = page.find_tables()
            for table in tables:
                cats = _parse_extracted_table(table.extract())
                if cats:
                    categories.extend(cats)
        except Exception:
            pass

    if categories:
        doc.close()
        return categories

    full_text = '\n'.join(page.get_text() for page in doc)
    doc.close()

    lines = [line.strip() for line in full_text.split('\n') if line.strip()]
    return _extract_sections_from_lines(lines)
